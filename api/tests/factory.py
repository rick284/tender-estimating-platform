"""Build a minimal but valid synthetic Master Estimate Workbook for tests (no client data)."""
from __future__ import annotations
import io
import openpyxl


def make_workbook(project="Test Depot Civil Works", rev="Rev1", qty=100, rate=100.0, oh=0.10, margin=0.15,
                  coast=50.0) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ["Cover & Instructions", "Resource Rate Library", "QS Takeoff", "Resource Build-Up", "Commercial Inputs",
                 "Weather Risk", "Consolidated Package", "Tender Build-Up", "Risk Truth", "Rate Verification"]:
        wb.create_sheet(name)
    cv = wb["Cover & Instructions"]
    cv["A2"] = f"{rev} test snapshot | synthetic"
    for r, (k, v) in enumerate([("Project", project), ("Location", "Nowhere, NZ"), ("Issue date", "01-Jan-2026"),
                                ("Release status", "WORKING")], start=10):
        cv.cell(r, 1, k); cv.cell(r, 2, v)

    lib = wb["Resource Rate Library"]
    lib.cell(5, 1, "PL-EXC"); lib.cell(5, 2, "Plant"); lib.cell(5, 3, "14 t excavator wet hire"); lib.cell(5, 5, "hr")
    lib.cell(5, 7, "No"); lib.cell(5, 8, "No"); lib.cell(5, 9, rate)

    q = wb["QS Takeoff"]
    q.cell(3, 1, "Earthworks"); q.cell(3, 2, "EW01"); q.cell(3, 3, "Strip"); q.cell(3, 4, "m2"); q.cell(3, 5, qty)
    q.cell(3, 8, "MEASURED"); q.cell(3, 10, "AMBER - test")

    b = wb["Resource Build-Up"]
    b["B7"] = "EW01"; b["N7"] = "Time x Count"; b["O7"] = "Plant"; b["P7"] = "PL-EXC"; b["I7"] = 50; b["R7"] = 1
    b["U7"] = 1; b["V7"] = 0; b["H7"] = "Strip"; b["AH7"] = "Yes"; b["AK7"] = "No"

    ci = wb["Commercial Inputs"]
    refs = [(f"CA{n:02d}", "PC Sum", "No") for n in range(1, 26)]
    for i, (ref, method, inc) in enumerate(refs):
        r = 5 + i
        ci.cell(r, 1, ref); ci.cell(r, 2, ref); ci.cell(r, 3, method); ci.cell(r, 6, inc)
    ci["C21"] = "Markup %"; ci["D21"] = oh; ci["F21"] = "Yes"
    ci["C22"] = "Margin %"; ci["D22"] = margin; ci["F22"] = "Yes"

    w = wb["Weather Risk"]
    w["B6"] = "Auckland"; w["B7"] = "2026-10-01"; w["B8"] = 55; w["B9"] = 5
    w["B13"] = "='Resource Build-Up'!Y7*8"; w["B14"] = "='Resource Build-Up'!Y7+'Resource Build-Up'!Y7"
    w["B16"] = 0; w["B19"] = "tool"; w["B20"] = "station"; w["B23"] = 0; w["B24"] = 0; w["B26"] = 0; w["B27"] = 0
    w["B28"] = 0; w["B29"] = "No"; w["B30"] = "AMBER - test"; w["B31"] = "current"; w["B38"] = 0

    cp = wb["Consolidated Package"]
    cp["B6"] = coast; cp["B7"] = 0; cp["B8"] = 0; cp["B10"] = 0; cp["B13"] = 0
    # labels and the "included in siteworks direct?" column as in the real template: the importer reads packages by what the rows say
    for r, (label, cls) in {6: ("Coast Concrete foundation/floor subcontract", "QUOTED"), 7: ("Allied public stormwater/detention base", "QUOTED BASE"),
                            8: ("Allied public setout/as-built/signoff additional", "CONFIRMED ADDITIONAL"), 10: ("Allied private drainage total", "QUOTED / ALLOCATED")}.items():
        cp.cell(r, 1, label); cp.cell(r, 3, cls); cp.cell(r, 4, "No")
    cp["A13"] = "Confirmed Siteworks/Coast overlap after clarification"
    cp["A12"] = "GROSS CURRENT KNOWN PACKAGE"

    base = qty / 50 * rate
    sub = base
    ohv = sub * oh
    profit = (sub + ohv) / (1 - margin) - (sub + ohv)
    tb = wb["Tender Build-Up"]
    tb["B20"] = base; tb["B39"] = 0; tb["B40"] = sub; tb["B42"] = ohv; tb["B43"] = profit; tb["B47"] = sub + ohv + profit
    gross = sub + coast
    cp["B12"] = gross
    cp["F56"] = (gross * (1 + oh)) / (1 - margin)

    rt = wb["Risk Truth"]
    rt["B12"] = 0; rt["B13"] = 0; rt["B15"] = "PASS"; rt["B16"] = 0; rt["B17"] = 0; rt["B18"] = "OK"; rt["B19"] = "OK"; rt["B21"] = "YES"
    rv = wb["Rate Verification"]
    rv["B4"] = 1; rv["B5"] = 0; rv["F4"] = 0; rv["C4"] = base; rv["C5"] = 0; rv["G4"] = 0

    import datetime as _dt
    wb.properties.created = wb.properties.modified = _dt.datetime(2026, 1, 1)      # fixed, so two builds of the same workbook are byte-identical (the duplicate check hashes the file)
    buf = io.BytesIO(); wb.save(buf)
    import zipfile                                                                  # zip entries carry a timestamp too: rewrite them with a fixed one
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED
            data = zin.read(info.filename)
            if info.filename == "docProps/core.xml":                                # openpyxl stamps "modified" with the save time
                import re
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            zout.writestr(zi, data)
    return out.getvalue()
