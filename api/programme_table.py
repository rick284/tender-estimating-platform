"""A programme given as a table (the 'MS Project Import' CSV / sheet): ID, Outline Level, WBS, Task Name, Duration, Predecessors, Start, Finish, Resource Names, Notes.
Pure functions: rows -> Gantt tasks (the same task shape the MS Project reader produces). Column names are matched case-insensitively; only Task Name is essential."""
from __future__ import annotations
import csv
import io
import re
from typing import Any

SYNONYMS = {"id": ("id", "uid", "task id", "unique id"), "level": ("outline level", "level"), "wbs": ("wbs", "code", "qs code"), "name": ("task name", "name", "task"),
            "duration": ("duration",), "preds": ("predecessors", "predecessor"), "start": ("start",), "finish": ("finish",), "resources": ("resource names", "resources"),
            "notes": ("notes", "note"), "category": ("category",)}


class BadTable(ValueError):
    pass


def duration_days(text: Any) -> int | None:
    t = str(text if text is not None else "").strip().lower()
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(days?|d|wks?|weeks?|w|hrs?|hours?|h|mons?|months?)?", t)
    if not m:
        return None
    n, unit = float(m.group(1)), (m.group(2) or "d")
    n = n * 5 if unit.startswith("w") else n * 20 if unit.startswith("mon") else n / 8 if unit.startswith("h") else n
    return int(-(-n // 1))


def _cols(header: list[str]) -> dict[str, int]:
    low = [h.strip().lower() for h in header]
    out = {}
    for key, names in SYNONYMS.items():
        for n in names:
            if n in low:
                out[key] = low.index(n)
                break
    if "name" not in out:
        raise BadTable("The table needs a 'Task Name' column. Use the MS Project Import layout: ID, Outline Level, WBS, Task Name, Duration, Predecessors.")
    return out


def read_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def read_xlsx(data: bytes) -> list[list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    name = next((n for n in wb.sheetnames if n.lower().replace("-", " ") == "ms project import"), wb.sheetnames[0])
    return [["" if c is None else str(c) for c in row] for row in wb[name].iter_rows(values_only=True)]


def tasks_from_table(rows: list[list[str]]) -> list[dict[str, Any]]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if len(rows) < 2:
        raise BadTable("The table has no tasks.")
    c = _cols(rows[0])
    body = rows[1:]
    get = lambda r, k: (r[c[k]].strip() if k in c and c[k] < len(r) and r[c[k]] is not None else "")     # noqa: E731
    tasks: list[dict[str, Any]] = []
    uid_of: dict[str, int] = {}
    for i, r in enumerate(body, 1):
        name = get(r, "name")
        if not name:
            continue
        uid = i
        uid_of[get(r, "id") or str(i)] = uid
        lvl = get(r, "level")
        level = int(lvl) if lvl.isdigit() and int(lvl) > 0 else 1
        days = duration_days(get(r, "duration"))
        t = {"uid": uid, "name": name[:240], "wbs": get(r, "wbs"), "outline_level": level, "duration": {"value": days if days is not None else 1, "unit": "D"}, "percent_complete": 0,
             "predecessors": [], "resources": [], "resource_assignments": [], "notes": get(r, "notes")[:1000], "is_active": True, "source_ref": get(r, "wbs"),
             "_preds": get(r, "preds"), "_category": get(r, "category").lower()}
        if days == 0:
            t["is_milestone"] = True
        names = [n.strip() for n in re.split(r"[;,]", get(r, "resources")) if n.strip()]
        t["resources"] = names
        t["resource_assignments"] = [{"resource_id": None, "name": n, "type": "", "unit": None, "units": 1.0} for n in names]
        tasks.append(t)
    if not tasks:
        raise BadTable("No task names were found in the table.")
    for t in tasks:
        for tok in re.split(r"[;,\s]+", t.pop("_preds")):
            m = re.fullmatch(r"(\d+)(FS|SS|FF|SF)?(?:([+-]\d+)\s*(?:d|days?)?)?", tok.strip(), re.I)
            if m and m.group(1) in uid_of and uid_of[m.group(1)] != t["uid"]:
                t["predecessors"].append({"uid": uid_of[m.group(1)], "type": (m.group(2) or "FS").upper(), "lag_days": int(m.group(3) or 0)})
        t.pop("_category")
    return tasks
