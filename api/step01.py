"""Step 01 outputs from the senior estimator, read from Word files (pure, no network): the Scope Confirmation (scope rows + written-answer
interface questions) and the Estimating Proposal (hours and fee window). Nothing here prices or decides anything: the rows are offered to the
estimator, who adds them to the Scope Coverage draft and approves."""
from __future__ import annotations
import io
import re
from typing import Any

import docx

PROPOSALS = ("included separately", "included / interface", "included; detail to confirm", "included", "subcontract quote", "to confirm", "by others / confirm", "by others", "excluded")


def _cell_text(c) -> str:
    return re.sub(r"[ \t]+", " ", "\n".join(p.text for p in c.paragraphs)).strip()


def _tables(data: bytes) -> list[list[list[str]]]:
    d = docx.Document(io.BytesIO(data))
    return [[[_cell_text(c) for c in r.cells] for r in t.rows] for t in d.tables]


def _col(head: list[str], *words: str) -> int | None:
    for i, h in enumerate(head):
        if any(w in h.lower() for w in words):
            return i
    return None


def owner_treatment(proposal: str) -> tuple[str, str]:
    """The Scope Coverage owner and treatment a proposal word suggests. Nothing is PRICED: no cost exists at Step 01."""
    p = proposal.lower()
    if "excluded" in p:
        return "OTHERS", "EXCLUDED"
    if "by others" in p:
        return "OTHERS", "BY OTHERS"
    if "subcontract" in p:
        return "SUBCONTRACT", "UNRESOLVED"
    if "interface" in p:
        return "SHARED", "UNRESOLVED"
    if "to confirm" in p or "detail to confirm" in p:
        return "TBC", "UNRESOLVED"
    return "OWN WORKS", "UNRESOLVED"


def parse_scope_confirmation(data: bytes) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    questions: list[dict[str, str]] = []
    for t in _tables(data):
        if not t:
            continue
        head = t[0]
        i_ref, i_item, i_bound, i_prop, i_src = _col(head, "ref"), _col(head, "work item", "item"), _col(head, "boundar", "proposed work"), _col(head, "our proposal", "proposal"), _col(head, "source")
        if i_ref is not None and i_item is not None and i_prop is not None:
            for r in t[1:]:
                ref, item = r[i_ref], r[i_item]
                m = re.match(r"^(\d{1,3})\b\s*(.*)$", ref)
                if not m:
                    continue
                if not item and m.group(2):
                    item = m.group(2)
                prop = re.sub(r"\[.\]\s*(Confirm|Exclude|Query)", "", r[i_prop]).strip()
                rows.append({"ref": m.group(1).zfill(2), "title": re.sub(r"\s+", " ", item).strip(), "detail": re.sub(r"\s+", " ", r[i_bound] if i_bound is not None else "").strip(),
                             "proposal": re.sub(r"\s+", " ", prop), "source": re.sub(r"\s+", " ", r[i_src] if i_src is not None else "").strip()})
            continue
        i_q, i_what, i_ans = _col(head, "ref"), _col(head, "interface", "decision"), _col(head, "written answer", "answer")
        if i_q is not None and i_what is not None and i_ans is not None:
            for r in t[1:]:
                if re.match(r"^Q\d+$", r[i_q].strip()):
                    questions.append({"ref": r[i_q].strip(), "title": re.sub(r"\s+", " ", r[i_what]).strip(), "question": re.sub(r"\s+", " ", r[i_ans]).strip()})
    for r in rows:
        r["owner"], r["treatment"] = owner_treatment(r["proposal"])
    return {"rows": rows, "questions": questions}


def _money(s: str) -> float | None:
    m = re.search(r"\$?\s*([0-9][0-9,]*(?:\.\d+)?)", s or "")
    return float(m.group(1).replace(",", "")) if m else None


def parse_proposal(data: bytes) -> dict[str, Any]:
    """Hours and fee window from the proposal's table (label, hours, rate, fee) plus its text."""
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    out: dict[str, Any] = {"cases": [], "rate": None, "text": text[:6000]}
    for t in _tables(data):
        if not t:
            continue
        head = [h.lower() for h in t[0]]
        ih, ir, ifee = _col(head, "hours"), _col(head, "rate"), _col(head, "fee")
        if ih is None or ifee is None:
            continue
        for r in t[1:]:
            h, fee = _money(r[ih]), _money(r[ifee])
            if h is None or fee is None:
                continue
            rate = _money(r[ir]) if ir is not None else None
            out["cases"].append({"label": r[0], "hours": h, "rate": rate, "fee": fee})
            out["rate"] = out["rate"] or rate
    hrs = sorted(c["hours"] for c in out["cases"])
    if hrs:
        out["low_hours"], out["high_hours"] = hrs[0], hrs[-1]
        out["low_fee"], out["high_fee"] = min(c["fee"] for c in out["cases"]), max(c["fee"] for c in out["cases"])
    m = re.search(r"\bRev\s?([A-Z])\b", text)
    out["revision"] = m.group(1) if m else ""
    return out


FILE_KINDS = (("scope", re.compile(r"scope[_ ]?confirmation.*\.docx$", re.I)), ("proposal", re.compile(r"estimating[_ ]?proposal.*\.docx$", re.I)),
              ("card", re.compile(r"(decision[_ ]?card|effort).*\.md$", re.I)))


def pick_files(files: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Newest matching file of each kind (highest Rev letter, then most recently modified)."""
    best: dict[str, dict[str, Any]] = {}
    for f in files:
        for kind, rx in FILE_KINDS:
            if rx.search(f["name"]):
                rev = re.search(r"Rev\s?_?([A-Z]|\d+)", f["name"], re.I)
                key = (rev.group(1).upper() if rev else "", f.get("modified", ""))
                if kind not in best or key > best[kind]["_key"]:
                    best[kind] = {**f, "_key": key}
    return {k: {x: y for x, y in v.items() if x != "_key"} for k, v in best.items()}
