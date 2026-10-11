"""Turn a downloaded document into page-by-page text. Pure functions: bytes in, text out, no network or database.
Every piece of text keeps its page number so a quantity can later cite 'file, revision, page'."""
from __future__ import annotations
import csv
import io
import zipfile
from typing import Any

MAX_PAGES = 400
MAX_CHARS_PER_PAGE = 60_000
SCAN_THRESHOLD = 20          # fewer characters than this across a PDF page = probably a scanned image


def _clip(t: str) -> str:
    t = "\n".join(line.rstrip() for line in t.replace("\x00", "").splitlines())
    return t[:MAX_CHARS_PER_PAGE]


def _pdf(data: bytes) -> dict[str, Any]:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    if r.is_encrypted:
        try:
            if not r.decrypt(""):
                return {"pages": [], "method": "pdf", "warnings": ["The PDF is password protected."], "status": "failed"}
        except Exception:
            return {"pages": [], "method": "pdf", "warnings": ["The PDF is password protected."], "status": "failed"}
    pages, warn = [], []
    total = len(r.pages)
    for i, p in enumerate(r.pages[:MAX_PAGES], 1):
        try:
            pages.append({"n": i, "text": _clip(p.extract_text() or "")})
        except Exception:
            pages.append({"n": i, "text": ""}); warn.append(f"Page {i} could not be read.")
    if total > MAX_PAGES:
        warn.append(f"Only the first {MAX_PAGES} of {total} pages were read.")
    blank = [p["n"] for p in pages if len(p["text"].strip()) < SCAN_THRESHOLD]
    status = "ok"
    if pages and len(blank) == len(pages):
        status = "needs_ocr"; warn.append("No text found: this looks like a scanned or drawing-only PDF and needs OCR / drawing reading.")
    elif blank:
        status = "partial"; warn.append(f"{len(blank)} page(s) have no readable text (likely drawings or scans): {', '.join(map(str, blank[:12]))}{'…' if len(blank) > 12 else ''}.")
    return {"pages": pages, "method": "pdf", "warnings": warn, "status": status}


def _docx(data: bytes) -> dict[str, Any]:
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return {"pages": [{"n": 1, "text": _clip("\n".join(lines))}], "method": "docx", "warnings": ["Word files have no fixed pages; all text is shown as page 1."], "status": "ok"}


def _xlsx(data: bytes) -> dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets[:60], 1):
        rows = []
        for row in ws.iter_rows(values_only=True, max_row=3000):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        pages.append({"n": i, "text": _clip(f"[Sheet: {ws.title}]\n" + "\n".join(rows))})
    return {"pages": pages, "method": "xlsx", "warnings": ["Each sheet is one 'page'. Values are as last saved (formulas are not recalculated)."], "status": "ok"}


def _text(data: bytes, csv_mode: bool) -> dict[str, Any]:
    t = data.decode("utf-8", errors="replace")
    if csv_mode:
        t = "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(t)))
    return {"pages": [{"n": 1, "text": _clip(t)}], "method": "text", "warnings": [], "status": "ok"}


def extract(name: str, mime: str, data: bytes) -> dict[str, Any]:
    """Returns {pages:[{n,text}], method, warnings:[...], status: ok|partial|needs_ocr|unsupported|failed}."""
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    try:
        if ext == "pdf" or mime == "application/pdf":
            return _pdf(data)
        if ext == "docx":
            return _docx(data)
        if ext in ("xlsx", "xlsm"):
            return _xlsx(data)
        if ext in ("md", "txt", "csv", "tsv", "json") or mime.startswith("text/"):
            return _text(data, csv_mode=ext in ("csv", "tsv"))
        if ext == "zip" and zipfile.is_zipfile(io.BytesIO(data)):
            return {"pages": [], "method": "zip", "warnings": ["Zip archives are not opened."], "status": "unsupported"}
    except Exception as e:                                   # a damaged file must never break the batch
        return {"pages": [], "method": ext or "unknown", "warnings": [f"Could not read this file ({type(e).__name__})."], "status": "failed"}
    return {"pages": [], "method": ext or "unknown", "warnings": [f"{ext.upper() or 'This'} files are not supported yet (drawings and images need the drawing reader)."], "status": "unsupported"}
