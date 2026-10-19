"""Step 02: the estimating production programme (the plan for the ESTIMATING WORK itself, not the construction programme).

A list of tasks, each with an owner, low/high hours, what drives the hours, what it follows, what it delivers and how it is accepted, grouped in stages.
Pure functions only: read a ChatGPT task-plan workbook (tolerant of heading wording), check a plan, add it up against the Step 01 fee window and against the hours
actually tracked. Nothing here prices construction and nothing is invented: a task with no hours is flagged, not guessed.
"""
from __future__ import annotations
import io
import re
from typing import Any

FIELDS = ("stage", "name", "owner", "low", "high", "driver", "predecessor", "deliverable", "acceptance", "scope_refs")
ALIASES: dict[str, tuple[str, ...]] = {
    "ref": ("id", "task id", "ref", "no", "task no"),
    "stage": ("stage", "phase", "group", "work package", "package"),
    "name": ("task", "task name", "activity", "description", "work item"),
    "owner": ("owner", "responsible", "who", "assigned"),
    "low": ("low hours", "low h", "hours low", "low", "min hours", "min"),
    "high": ("high hours", "high h", "hours high", "high", "max hours", "max"),
    "hours": ("hours", "planned hours", "est hours", "estimated hours", "h"),
    "driver": ("source / driver", "source/driver", "driver", "source", "basis"),
    "predecessor": ("predecessor", "predecessors", "follows", "depends on", "after"),
    "deliverable": ("deliverable", "output", "outputs"),
    "acceptance": ("acceptance test", "acceptance", "acceptance criteria", "done when", "test"),
    "scope_refs": ("scope rows", "scope refs", "scope", "scope ref"),
}


def _n(s: Any) -> str:
    return re.sub(r"[^a-z0-9/ ]+", " ", str(s or "").lower()).strip()


def _num(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        m = re.search(r"\d+(?:\.\d+)?", str(v))
        return float(m.group()) if m else None


def _map_header(cells: list[Any]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, c in enumerate(cells):
        h = _n(c)
        if not h:
            continue
        for field, names in ALIASES.items():
            if field not in out and h in names:
                out[field] = i
                break
    return out


def parse_task_plan(data: bytes) -> list[dict[str, Any]]:
    """Tasks from the workbook's task-plan sheet: the first sheet with a header row that has a task/activity column and an hours column."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(max_row=400, values_only=True)]
        for hi, r in enumerate(rows[:30]):
            m = _map_header(r)
            if "name" in m and ({"low", "high"} & m.keys() or "hours" in m):
                tasks = []
                for r2 in rows[hi + 1:]:
                    g = lambda k: (r2[m[k]] if k in m and m[k] < len(r2) else None)
                    name = str(g("name") or "").strip()
                    if not name or _n(name).startswith("total"):
                        continue
                    lo, hi_ = _num(g("low")), _num(g("high"))
                    single = _num(g("hours"))
                    if lo is None and hi_ is None:
                        lo = hi_ = single
                    elif lo is None or hi_ is None:
                        lo = hi_ = lo if lo is not None else hi_
                    tasks.append({"ref": str(g("ref") or "").strip(), "stage": str(g("stage") or "").strip(), "name": name[:200], "owner": str(g("owner") or "").strip()[:80],
                                  "low": lo, "high": hi_, "driver": str(g("driver") or "").strip()[:500], "predecessor": str(g("predecessor") or "").strip()[:100],
                                  "deliverable": str(g("deliverable") or "").strip()[:500], "acceptance": str(g("acceptance") or "").strip()[:500], "scope_refs": str(g("scope_refs") or "").strip()[:300]})
                if tasks:
                    return tasks
    raise ValueError("No task plan found: the workbook needs a sheet with a Task (or Activity) column and Hours (or Low / High hours) columns.")


def clean(tasks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Validated copy for saving. Hours must be numbers >= 0 and low <= high; actual hours are kept as typed."""
    out, seen = [], set()
    for i, t in enumerate(tasks[:500], 1):
        name = str(t.get("name") or "").strip()
        if not name:
            raise ValueError(f"Task {i} has no name.")
        lo, hi = t.get("low"), t.get("high")
        lo = None if lo in (None, "") else float(lo)
        hi = None if hi in (None, "") else float(hi)
        for v in (lo, hi):
            if v is not None and (v < 0 or v > 2000):
                raise ValueError(f"Task '{name}': hours must be between 0 and 2000.")
        if lo is not None and hi is not None and lo > hi:
            raise ValueError(f"Task '{name}': low hours are above high hours.")
        act = t.get("actual")
        act = None if act in (None, "") else float(act)
        if act is not None and act < 0:
            raise ValueError(f"Task '{name}': actual hours cannot be negative.")
        uid = str(t.get("id") or "") or f"P-{i:03d}"
        if uid in seen:
            raise ValueError(f"Two tasks share the id {uid}.")
        seen.add(uid)
        out.append({"id": uid, "ref": str(t.get("ref") or "")[:20], "stage": str(t.get("stage") or "")[:80], "name": name[:200], "owner": str(t.get("owner") or "")[:80], "low": lo, "high": hi,
                    "driver": str(t.get("driver") or "")[:500], "predecessor": str(t.get("predecessor") or "")[:100], "deliverable": str(t.get("deliverable") or "")[:500],
                    "acceptance": str(t.get("acceptance") or "")[:500], "scope_refs": str(t.get("scope_refs") or "")[:300], "actual": act, "done": bool(t.get("done"))})
    return out


def summarise(tasks: list[dict[str, Any]], window: dict[str, Any] | None, tracked_hours: float) -> dict[str, Any]:
    """Totals per stage, the reconciliation to the Step 01 window and plain-language checks. The platform only adds up and compares; it never fixes a plan."""
    lo = sum(t["low"] or 0 for t in tasks)
    hi = sum(t["high"] or 0 for t in tasks)
    stages: dict[str, dict[str, Any]] = {}
    for t in tasks:
        s = stages.setdefault(t["stage"] or "(no stage)", {"stage": t["stage"] or "(no stage)", "tasks": 0, "done": 0, "low": 0.0, "high": 0.0, "actual": 0.0})
        s["tasks"] += 1; s["done"] += 1 if t.get("done") else 0
        s["low"] += t["low"] or 0; s["high"] += t["high"] or 0; s["actual"] += t.get("actual") or 0
    issues: list[str] = []
    no_hours = [t["name"] for t in tasks if t["low"] is None and t["high"] is None]
    if no_hours:
        issues.append(f"{len(no_hours)} task(s) have no hours: {', '.join(no_hours[:4])}{'...' if len(no_hours) > 4 else ''}.")
    for k, label in (("owner", "an owner"), ("acceptance", "an acceptance test"), ("deliverable", "a deliverable")):
        n = sum(1 for t in tasks if not t[k])
        if n:
            issues.append(f"{n} task(s) have no {label.split(' ', 1)[1]}.")
    names = {t["ref"] for t in tasks if t["ref"]} | {t["id"] for t in tasks}
    bad = [t["name"] for t in tasks if t["predecessor"] and not all(p.strip() in names for p in re.split(r"[;,]", t["predecessor"]) if p.strip())]
    if bad:
        issues.append(f"{len(bad)} task(s) follow a task that is not in the list: {', '.join(bad[:3])}.")
    recon = None
    if window and window.get("low_hours") is not None and window.get("high_hours") is not None:
        wl, wh = float(window["low_hours"]), float(window["high_hours"])
        pos = "below" if hi < wl else "above" if lo > wh else "within"
        recon = {"low_hours": wl, "high_hours": wh, "position": pos, "benchmark": window.get("benchmark_hours"),
                 "note": {"within": "The plan's range overlaps the Step 01 window.", "below": "The plan is below the Step 01 window: the fee may be overstated, or tasks are missing.",
                          "above": "The plan is above the Step 01 window: the fee may be understated. Explain the variance with evidence before the go-ahead."}[pos]}
    actual = sum(t.get("actual") or 0 for t in tasks)
    return {"low": round(lo, 2), "high": round(hi, 2), "mid": round((lo + hi) / 2, 2), "actual_task_hours": round(actual, 2), "tracked_hours": round(tracked_hours, 2),
            "stages": [{**s, "low": round(s["low"], 2), "high": round(s["high"], 2), "actual": round(s["actual"], 2)} for s in stages.values()],
            "done": sum(1 for t in tasks if t.get("done")), "count": len(tasks), "issues": issues, "reconciliation": recon}
