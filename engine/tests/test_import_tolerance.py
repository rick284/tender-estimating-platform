"""A workbook from an older or different job doesn't have every sheet or the same number of rows as Hobsonville. It must still load, say what was left out,
and never silently drop quantities."""
import io
import json
from pathlib import Path

import openpyxl
import pytest

from tender_engine.estimate import compute_estimate
from tender_engine.importers import load_master_workbook
from tender_engine.inputs import from_json, to_json

PRIV = Path(__file__).resolve().parents[2] / "private"
E, DRAIN = PRIV / "Rev34E.xlsx", PRIV / "Drainage_Rev3.xlsx"
needs_e = pytest.mark.skipif(not E.exists(), reason="private workbook not present")
needs_drain = pytest.mark.skipif(not DRAIN.exists(), reason="private workbook not present")


def without(sheets):
    wb = openpyxl.load_workbook(E)
    for s in sheets:
        del wb[s]
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@needs_drain
def test_older_job_without_a_consolidated_package_loads_and_reconciles_to_the_cent():
    """329 Hobsonville Drainage Rev3: no Consolidated Package sheet, weather links not formulas, and two QS items (TC01, TC02: $11,000) below row 40."""
    m = load_master_workbook(DRAIN)
    e = compute_estimate(m)
    assert {q.code for q in m.qs} >= {"SW01", "SW08", "TC01", "TC02"}                       # items below the template's own rows are read
    assert round(e["tender"]["base_direct"], 2) == 218077.81 and round(e["tender"]["total_excl_gst"], 2) == 232838.98
    assert e["reconciled"] is True and all(r["ok"] for r in e["reconciliation"])             # every figure the workbook shows is reproduced
    assert round(e["summary"]["all_in"], 2) == round(e["tender"]["total_excl_gst"], 2)        # no subcontract packages: the all-in equals the siteworks tender
    notes = " ".join(e["import_notes"])
    assert "No Consolidated Package" in notes and "weather costs are not linked" in notes
    assert e["import_notes"] == to_json(from_json(json.loads(json.dumps(to_json(m)))))["notes"]       # the notes are saved with the revision


@needs_e
@pytest.mark.parametrize("gone,expect_note", [
    (["Consolidated Package"], "No Consolidated Package"), (["Risk Truth"], "No Risk Truth"), (["Rate Verification"], "No Rate Verification"),
    (["Weather Risk"], "No Weather Risk"), (["Tender Build-Up"], "No Tender Build-Up"), (["Cover & Instructions"], None),
    (["Consolidated Package", "Risk Truth", "Weather Risk", "Rate Verification"], "No Consolidated Package"),
])
def test_any_optional_sheet_can_be_missing(gone, expect_note):
    m = load_master_workbook(without(gone))
    e = compute_estimate(m)                                                                # the whole estimate still calculates
    assert len(m.qs) == 30 and round(e["tender"]["base_direct"], 2) == 435521.88            # the quantities are untouched
    if expect_note:
        assert any(expect_note in n for n in e["import_notes"])
    if "Tender Build-Up" in gone:
        assert e["reconciled"] is None and e["reconciliation"] == []                          # nothing to compare with: no false mismatch
    if "Consolidated Package" in gone:
        assert e["consolidated"]["overlays_total"] == 0


@needs_e
def test_the_real_workbook_is_unchanged_by_all_this():
    m = load_master_workbook(E)
    e = compute_estimate(m)
    assert round(e["summary"]["all_in"], 2) == 1226916.81 and e["reconciled"] is True and e["import_notes"] == []
    assert [o.name for o in m.overlays] == ["Coast Concrete foundation / floor", "Allied public stormwater / detention base",
                                           "Allied public setout / as-built / signoff", "Allied private drainage"]    # names the schedules and saved work rely on


def test_an_essential_sheet_missing_says_which_and_what_it_holds():
    wb = openpyxl.Workbook(); wb.active.title = "Cover & Instructions"; wb.create_sheet("QS Takeoff")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    with pytest.raises(ValueError) as ei:
        load_master_workbook(buf)
    msg = str(ei.value)
    assert "Resource Rate Library (the rates)" in msg and "Resource Build-Up (the cost build-up rows)" in msg and "Commercial Inputs" in msg and "QS Takeoff" not in msg.split("missing:")[1].split(".")[0]
