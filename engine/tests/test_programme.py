import datetime as dt
from pathlib import Path
import pytest

from tender_engine import D
from tender_engine.blank import blank_from_template
from tender_engine.edits import EditError, apply_edits
from tender_engine.estimate import compute_estimate
from tender_engine.importers import load_master_workbook
from tender_engine.inputs import from_json, to_json
from tender_engine.programme import compute_programme, derived_days, new_programme, working_date

PRIVATE = Path(__file__).resolve().parents[2] / "private" / "Rev34E.xlsx"


@pytest.fixture(scope="module")
def m():
    return load_master_workbook(PRIVATE)


def T(code, preds=(), days=None, ham=False, qs=""):
    return {"id": f"P-{code}xxx", "code": code, "name": code, "qs_code": qs, "override_days": None if days is None else str(days), "predecessors": list(preds), "hammock": ham, "basis": "", "note": ""}


def test_derived_duration_rounds_up_to_working_days():
    assert derived_days(D("26.3125")) == 4 and derived_days(D("16")) == 2 and derived_days(D("16.01")) == 3 and derived_days(D(0)) is None and derived_days(None) is None
    assert derived_days(D("20"), 10) == 2


def test_working_dates_skip_weekends_and_roll_a_weekend_start():
    assert working_date(dt.date(2026, 10, 1), 0) == dt.date(2026, 10, 1)
    assert working_date(dt.date(2026, 10, 1), 3) == dt.date(2026, 10, 6)          # Thu +3 working days = Tue
    assert working_date(dt.date(2026, 10, 3), 0) == dt.date(2026, 10, 5)          # Saturday start rolls to Monday
    assert working_date(dt.date(2026, 10, 1), 3, 6) == dt.date(2026, 10, 5)       # six-day week: Sat counts


def test_critical_path_float_hammock_and_problems():
    p = new_programme("2026-10-01")
    p["tasks"] = [T("A", days=3), T("B", ["A"], 4), T("C", ["A"], 1), T("D", ["B", "C"], 2), T("PG", days=9, ham=True)]
    r = compute_programme(p, {}, {})
    assert r["duration"] == 9 and r["critical_path"] == ["A", "B", "D"] and r["finish"] == "2026-10-13" and not r["problems"]
    f = {t["code"]: t["float"] for t in r["tasks"]}
    assert f == {"A": 0, "B": 0, "C": 3, "D": 0, "PG": 0}
    p["tasks"][4]["override_days"] = "12"
    assert any("Preliminaries run 12 days" in x for x in compute_programme(p, {}, {})["problems"])
    p["tasks"][0]["predecessors"] = ["Z"]
    assert any("isn't in the programme" in x for x in compute_programme(p, {}, {})["problems"])
    q = new_programme(""); q["tasks"] = [T("A", ["B"], 1), T("B", ["A"], 1)]
    assert any("loop" in x for x in compute_programme(q, {}, {})["problems"])


def test_duration_comes_from_the_slowest_resource_row_unless_overridden():
    p = new_programme(""); p["tasks"] = [T("EW01", days=None, qs="EW01"), T("EW02", days=7, qs="EW02")]
    info = {"EW01": {"section": "Earthworks", "description": "d", "amount": 1.0}, "EW02": {"section": "Earthworks", "description": "d", "amount": 1.0}}
    r = compute_programme(p, {"EW01": D("26.3125"), "EW02": D("100")}, info)
    assert [t["duration"] for t in r["tasks"]] == [4, 7] and [t["derived"] for t in r["tasks"]] == [4, 13]


@pytest.mark.skipif(not PRIVATE.exists(), reason="private workbook not present")
class TestHobsonville:

    def test_programme_reproduces_the_workbook(self, m):
        p = compute_estimate(m)["programme"]
        assert p["set_up"] and p["duration"] == 55 and p["finish"] == "2026-12-16" and not p["problems"]
        assert p["critical_path"] == ["ESC01", "EW01", "PRG-RF01", "EW03", "ES01", "PRG-PROOF", "PV01", "PRG-CURE", "PV05", "RH01"]     # as stated in report 03
        by = {t["code"]: t for t in p["tasks"]}
        assert by["PG01"]["hammock"] and by["PG01"]["duration"] == 55 and not by["PG01"]["critical"]
        assert by["PV01"]["derived"] == 20 and by["PV01"]["duration"] == 10 and by["ESC01"]["derived"] == 2 and by["ESC01"]["duration"] == 3     # override beats derived
        assert by["EW06"]["duration"] == 0 and by["EW09"]["duration"] == 9 and by["DRN-REF01"]["duration"] == 21 and by["PRG-RF01"]["qs_code"] == ""

    def test_derived_durations_match_the_workbook_for_every_costed_task(self, m):
        import warnings; warnings.filterwarnings("ignore")
        import openpyxl
        v = openpyxl.load_workbook(PRIVATE, data_only=True)["Programme Export"]
        wb = {str(v.cell(r, 4).value): v.cell(r, 11).value for r in range(7, v.max_row + 1) if isinstance(v.cell(r, 4).value, str) and v.cell(r, 13).value is not None}
        tasks = {t["code"]: t for t in compute_estimate(m)["programme"]["tasks"]}
        bad = [(c, tasks[c]["derived"], k) for c, k in wb.items() if c in tasks and tasks[c]["qs_code"] and (tasks[c]["derived"] or 0) != (k or 0)]
        assert bad == []

    def test_setup_edit_and_round_trip_on_a_new_tender(self, m):
        blank = blank_from_template(m, "x", {})
        e0 = compute_estimate(blank)["programme"]
        assert e0["set_up"] is False and e0["missing"] == []
        edits = [{"target": "qs", "op": "add", "key": "EW01", "value": {"section": "Earthworks", "description": "Strip", "unit": "m2", "quantity": "1000", "basis": "MEASURED"}},
                 {"target": "qs", "op": "add", "key": "EW02", "value": {"section": "Earthworks", "description": "Fill", "unit": "m3", "quantity": "500", "basis": "MEASURED"}},
                 {"target": "row", "op": "add", "key": "N-prog0001", "value": {"qs_code": "EW01", "mode": "Lump Sum / Quote", "resource_code": "", "task": "Strip"}},
                 {"target": "row", "key": "N-prog0001", "field": "manual_qty", "value": "1"}, {"target": "row", "key": "N-prog0001", "field": "job_rate_override", "value": "1000"},
                 {"target": "programme", "op": "add", "key": "setup", "value": {"start": "2026-10-01"}},
                 {"target": "programme_task", "key": "P-EW01", "field": "override_days", "value": "4"},
                 {"target": "programme_task", "key": "P-EW02", "field": "override_days", "value": "6"},
                 {"target": "programme_task", "key": "P-EW02", "field": "predecessors", "value": "EW01"}]
        m2, ch = apply_edits(blank, edits)
        p = compute_estimate(m2)["programme"]
        assert p["duration"] == 10 and p["critical_path"] == ["EW01", "EW02"] and p["finish"] == "2026-10-14" and "programme" in {c["target"] for c in ch}
        assert compute_estimate(from_json(to_json(m2)))["programme"] == p
        m3, ch3 = apply_edits(m2, [{"target": "qs", "op": "delete", "key": "EW01"}])          # tasks of a deleted QS item go, with links to them
        t3 = {t["code"]: t for t in compute_estimate(m3)["programme"]["tasks"]}
        assert "EW01" not in t3 and t3["EW02"]["predecessors"] == []
        assert apply_edits(blank, edits[:2] + [{"target": "programme", "op": "add", "key": "setup", "value": {}}, {"target": "programme", "op": "delete", "key": "setup"}])[1][-1]["target"] == "qs"

    @pytest.mark.parametrize("edit,msg", [
        ({"target": "programme_task", "op": "add", "key": "P-bad00001", "value": {"code": "X1", "name": "x"}}, "Set up the programme first"),
        ({"target": "programme", "key": "config", "field": "start", "value": "2026-10-01"}, "Set up the programme first"),
        ({"target": "programme", "op": "delete", "key": "setup"}, "isn't set up"),
    ])
    def test_programme_validation_before_setup(self, m, edit, msg):
        with pytest.raises(EditError, match=msg):
            apply_edits(blank_from_template(m, "x", {}), [edit])

    @pytest.mark.parametrize("edit,msg", [
        ({"target": "programme_task", "op": "add", "key": "bad", "value": {"code": "X1", "name": "x"}}, "identity"),
        ({"target": "programme_task", "op": "add", "key": "P-new00001", "value": {"code": "ESC01", "name": "x"}}, "already used"),
        ({"target": "programme_task", "op": "add", "key": "P-new00001", "value": {"code": "bad code!", "name": "x"}}, "letters, numbers"),
        ({"target": "programme_task", "op": "add", "key": "P-new00001", "value": {"code": "X1"}}, "needs a name"),
        ({"target": "programme_task", "key": "P-00001", "field": "override_days", "value": "2.5"}, "whole number"),
        ({"target": "programme_task", "key": "P-00001", "field": "override_days", "value": "-1"}, "0 to 3000"),
        ({"target": "programme_task", "key": "P-00001", "field": "predecessors", "value": "PG01"}, "follow itself"),
        ({"target": "programme_task", "key": "P-zzzzzz", "field": "name", "value": "x"}, "Unknown programme task"),
        ({"target": "programme", "key": "setup", "op": "add", "value": {}}, "already set up"),
        ({"target": "programme", "key": "config", "field": "days_per_week", "value": "9"}, "between 1 and 7"),
        ({"target": "programme", "key": "config", "field": "start", "value": "soon"}, "valid date"),
    ])
    def test_programme_validation(self, m, edit, msg):
        with pytest.raises(EditError, match=msg):
            apply_edits(load_master_workbook(PRIVATE), [edit])

    def test_programme_workbook(self, m):
        openpyxl = pytest.importorskip("openpyxl")
        import io
        from tender_engine.outputs import OutputContext, programme_workbook
        est = compute_estimate(m)
        wb = openpyxl.load_workbook(io.BytesIO(programme_workbook(est, OutputContext("3Bears", "Rick", "329 Hobsonville Road", "Dave", "", "Rev34E", "12 September 2026"))))
        assert wb.sheetnames == ["Programme", "MS Project Import"]
        assert "ESC01 > EW01 > PRG-RF01" in wb["Programme"]["A3"].value
        imp = {r[2]: r for r in wb["MS Project Import"].iter_rows(min_row=2, values_only=True) if r[2]}
        assert imp["EW01"][4] == "4 days" and imp["EW01"][1] == 2 and imp["EW01"][7] == "2026-10-06"
        ids = {r[2]: r[0] for r in imp.values()}
        assert imp["EW01"][5] == str(ids["ESC01"])
        with pytest.raises(ValueError):
            programme_workbook(compute_estimate(blank_from_template(m, "x", {})), OutputContext())
