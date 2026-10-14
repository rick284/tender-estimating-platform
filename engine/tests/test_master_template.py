"""The clean master workbook template (instructions/Master_Estimate_Workbook_TEMPLATE.xlsx): a small job filled into it recalculates in LibreOffice, imports into the
platform and reconciles to the workbook's own totals. Skipped when LibreOffice is not installed."""
import datetime
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from tender_engine.estimate import compute_estimate
from tender_engine.importers import load_master_workbook

TEMPLATE = Path(__file__).resolve().parents[2] / "instructions" / "Master_Estimate_Workbook_TEMPLATE.xlsx"
SOFFICE = shutil.which("soffice")


def recalc(src: Path, tmp: Path) -> Path:
    out = tmp / "out"
    out.mkdir(exist_ok=True)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=240)
    return out / src.name


def fill_demo(path: Path) -> None:
    wb = openpyxl.load_workbook(TEMPLATE)
    qs = wb["QS Takeoff"]
    for r, (sec, code, d, u, q) in zip((3, 4, 5, 6), [("Preliminaries & General", "PG01", "Preliminaries", "LS", 1), ("Earthworks", "EW01", "Strip", "m2", 1000),
                                                      ("Earthworks", "EW02", "Imported fill", "m3", 300), ("Services", "SV01", "Specialist package", "LS", 1)]):
        qs[f"A{r}"], qs[f"B{r}"], qs[f"C{r}"], qs[f"D{r}"], qs[f"E{r}"], qs[f"I{r}"], qs[f"P{r}"] = sec, code, d, u, q, "DOC-001 sheet C1 Rev A p1", "AMBER"
    bu = wb["Resource Build-Up"]
    rows = [dict(B="PG01", H="Supervision", L=440, N="Time x Count", O="Labour", P="LAB-SUPERVISOR", R=1, U=1),
            dict(B="EW01", H="Strip workfront", I=80, J="m2/hr", N="Time x Count", O="Plant", P="PL-EXC-13", R=1, U=1, AK="Yes", AL="Earthworks"),
            dict(B="EW01", H="Strip labour", I=80, J="m2/hr", N="Time x Count", O="Labour", P="LAB-LABOURER", R=2, U=1, AK="Yes", AL="Earthworks"),
            dict(B="EW02", H="Import fill", N="QS Qty x Factor", O="Material", P="MAT-IMPFILL", R=1, V=0.05),
            dict(B="SV01", H="Quoted package", N="Lump Sum / Quote", O="Subcontract", P="SUB-SPECIALIST", S=1, X=5000)]
    for i, d in enumerate(rows):
        for col, v in d.items():
            bu[f"{col}{7 + i}"] = v
    wb["Commercial Inputs"]["E13"], wb["Commercial Inputs"]["F13"] = 3000, "Yes"
    cp = wb["Consolidated Package"]
    cp["A6"], cp["B6"], cp["C6"], cp["D6"] = "Package A", 100000, "QUOTED", "No"
    tb = wb["Tender Build-Up"]
    tb["A10"], tb["A11"], tb["A12"] = "Preliminaries & General", "Earthworks", "Services"
    wb["Cover & Instructions"]["B38"] = "Demo job"
    wr = wb["Weather Risk"]
    wr["B6"], wr["B7"], wr["B8"], wr["B9"] = "Auckland", datetime.datetime(2026, 11, 2), 55, 5
    sq = wb["Supplier Quote Register"]
    for col, v in zip("ABCDEFGHIJKL", ["Q-01", "Package A", "SV01", 1, "LS", 100000, "Scope", "Docs", "Supplier X", "REQUESTED", "Due Fri", "Ref"]):
        sq[f"{col}5"] = v
    rt = wb["Risk Truth"]
    for col, v in zip("ABCDEF", ["RT01", "Mass balance unknown", "D/E - Unknown extent / missing information", "High", "High", "Critical"]):
        rt[f"{col}25"] = v
    rt["K25"] = "RED"
    for col, v in {"A": "RT01-001", "B": "RT01", "C": "Fill volume", "D": "Quantity / Scope", "E": "EW02", "I": "RFI Hold", "P": "Primary", "Q": "Yes", "R": "CA01", "S": "No", "W": "RED", "Y": "No"}.items():
        rt[f"{col}41"] = v
    pe = wb["Programme Export"]
    pe["P8"], pe["P9"], pe["L7"] = "PG01", "EW01", 55
    wb.save(path)


@pytest.mark.skipif(not SOFFICE or not TEMPLATE.exists(), reason="needs LibreOffice and the committed template")
def test_a_job_filled_into_the_template_imports_and_reconciles(tmp_path):
    job = tmp_path / "job.xlsx"
    fill_demo(job)
    imp = load_master_workbook(recalc(job, tmp_path))
    est = compute_estimate(imp)
    assert [q.code for q in imp.qs] == ["PG01", "EW01", "EW02", "SV01"] and len(imp.rows) == 5
    s = est["summary"]
    assert round(s["base_direct"], 2) == 60420.00                      # 32,120 supervision + 2,825 strip + 20,475 fill + 5,000 package
    assert round(s["siteworks_tender"], 2) == 82072.94 and round(s["all_in"], 2) == 211484.71
    assert est["reconciled"] is True                                   # the platform's totals equal the workbook's own formulas
    assert imp.meta["project"] == "Demo job"
    assert [o.name for o in imp.overlays] == ["Package A"]
    assert [q["ref"] for q in imp.quotes] == ["Q-01"] and [e["ref"] for e in imp.risk_truth] == ["RT01-001"]
    assert [(t["code"], t["predecessors"]) for t in imp.programme["tasks"]][:3] == [("PG01", []), ("EW01", ["PG01"]), ("EW02", ["EW01"])]


@pytest.mark.skipif(not SOFFICE or not TEMPLATE.exists(), reason="needs LibreOffice and the committed template")
def test_ms_project_import_sheet_is_the_loadable_programme_table(tmp_path):
    """The template generates the MS Project Import CSV itself (compact rows, predecessors as task IDs); the platform reads that same table."""
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "api"))
    import programme_table
    job = tmp_path / "job.xlsx"
    fill_demo(job)
    done = recalc(job, tmp_path)
    wb = openpyxl.load_workbook(done, data_only=True)
    ms = wb["MS Project Import"]
    rows = [[("" if c.value is None else str(c.value)) for c in r][:17] for r in ms.iter_rows(min_row=1, max_row=6)]
    assert [r[2] for r in rows[1:5]] == ["PG01", "EW01", "EW02", "SV01"] and rows[1][4] == "55 days" and rows[2][4] == "2 days"
    assert rows[2][5] == "1" and rows[3][5] == "2"                                       # EW01 follows task 1 (PG01), EW02 follows task 2 (EW01)
    tasks = programme_table.tasks_from_table(programme_table.read_xlsx(done.read_bytes()))
    assert [t["name"] for t in tasks][:2] == ["Preliminaries", "Strip"] and tasks[1]["predecessors"] == [{"uid": 1, "type": "FS", "lag_days": 0}]


@pytest.mark.skipif(not SOFFICE or not TEMPLATE.exists(), reason="needs LibreOffice and the committed template")
def test_the_blank_template_itself_imports(tmp_path):
    imp = load_master_workbook(recalc(TEMPLATE, tmp_path))
    est = compute_estimate(imp)
    assert imp.qs == [] and est["summary"]["all_in"] == 0
