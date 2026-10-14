"""Build the clean Master Estimate Workbook template for external AI assistants.

Starts from a real job workbook (as a donor of styles, validations and generic formulas), strips every job-specific value, rewrites the job-bound sheets
(Cover, Consolidated Package, Tender Build-Up sections, Programme Export links) in generic form, extends the formula blocks, and adds the log sheets
(Email, RFI, Quote, Document Register, Decision, Contacts, Change Log) and a Template Guide.  The result is written to instructions/.

usage: python make_master_template.py DONOR.xlsx OUT.xlsx
"""
from __future__ import annotations
import copy
import re
import sys

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

VERSION = "Rev1 (14 Oct 2026)"
QS_LAST, BU_LAST, LIB_LAST, PE_LAST, RV_LAST = 202, 306, 200, 206, 306
KEEP = ["Cover & Instructions", "Consolidated Package", "Tender Build-Up", "Commercial Inputs", "QS Takeoff", "Resource Build-Up", "Resource Rate Library", "Rate Verification",
        "Programme Export", "Weather Risk", "Risk Truth", "Supplier Quote Register", "Source & Markup Register"]


def isf(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def put(ws, addr, val):
    c = ws[addr]
    if not isinstance(c, MergedCell):
        c.value = val


def wipe(ws, r1, r2, c1=1, c2=None, formulas=False):
    """Clear typed values (and formulas when asked) in a block. Merged cells are left alone."""
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2 or ws.max_column):
        for c in row:
            if isinstance(c, MergedCell) or c.value is None:
                continue
            if formulas or not isf(c.value):
                c.value = None


def unmerge_all(ws):
    for m in list(ws.merged_cells.ranges):
        ws.merged_cells.remove(m)
        for r, c in m.cells:
            if isinstance(ws._cells.get((r, c)), MergedCell):
                del ws._cells[(r, c)]


def style_like(dst, src):
    dst.font, dst.fill, dst.border = copy.copy(src.font), copy.copy(src.fill), copy.copy(src.border)
    dst.alignment, dst.number_format, dst.protection = copy.copy(src.alignment), src.number_format, copy.copy(src.protection)


def tr(formula: str, src: str, dst: str) -> str:
    return Translator(formula, origin=src).translate_formula(dst)


def build(donor: str, out: str) -> None:
    wb = openpyxl.load_workbook(donor)
    for n in list(wb.sheetnames):
        if n not in KEEP:
            del wb[n]
    # ------------------------------------------------------------------ library: keep every rate (company asset); extend lookups to row 200
    lib = wb["Resource Rate Library"]
    put(lib, "A1", "COMPANY RESOURCE RATE LIBRARY - MASTER ESTIMATE WORKBOOK")
    for r in range(82, LIB_LAST + 1):                               # more rows for new rates, same look as the existing rows
        for c in range(1, 21):
            style_like(lib.cell(r, c), lib.cell(80, c))
    # ------------------------------------------------------------------ QS Takeoff
    qs = wb["QS Takeoff"]
    g3 = qs["G3"].value
    tmpl = {c: qs[f"{c}60"].value for c in "FLMNOR"}
    tmpl["G"] = g3
    style_row = {c: qs.cell(60, c) for c in range(1, 21)}
    wipe(qs, 3, 137, 1, 20, formulas=True)
    for r in range(3, QS_LAST + 1):
        for c in range(1, 21):
            style_like(qs.cell(r, c), style_row[c]) if r > 60 else None
        for col, f in tmpl.items():
            src = "G3" if col == "G" else f"{col}60"
            qs[f"{col}{r}"].value = tr(f, src, f"{col}{r}")
    tr_ = QS_LAST + 1
    put(qs, "A1", "QS TAKEOFF - MASTER ESTIMATE WORKBOOK - RESOURCE-LINKED")
    put(qs, f"A{tr_}", "TOTAL DIRECT COST")
    put(qs, f"G{tr_}", f"=SUM(G3:G{QS_LAST})")
    for c in range(1, 21):
        style_like(qs.cell(tr_, c), qs.cell(2, c))
    put(qs, "L1", "Rate and Amount are linked from Resource Build-Up. One row per priced item; the section name goes on every row. Do not type over the grey formula columns.")
    qs.data_validations.dataValidation = []
    dv = DataValidation(type="custom", formula1="FALSE", allow_blank=True, showErrorMessage=True, errorTitle="Formula cell", error="This column is calculated. Change the Resource Build-Up instead.")
    dv.add(f"F3:G{QS_LAST}"); dv.add(f"L3:O{QS_LAST}"); dv.add(f"R3:R{QS_LAST}")
    qs.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1='"GREEN,AMBER,RED"', allow_blank=True)
    dv2.add(f"P3:P{QS_LAST}")
    qs.add_data_validation(dv2)
    # ------------------------------------------------------------------ Resource Build-Up
    bu = wb["Resource Build-Up"]
    formula_cols = [c for c in range(1, 41) if isf(bu.cell(BU_LAST, c).value)]
    for r in range(7, BU_LAST + 1):
        for c in range(1, 41):
            cell = bu.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            if c in formula_cols:
                cell.value = tr(bu.cell(BU_LAST, c).value, f"{bu.cell(BU_LAST, c).coordinate}", cell.coordinate) if (r != BU_LAST) else cell.value
            else:
                cell.value = None
    # ------------------------------------------------------------------ Resource Build-Up / Library / QS references: widen the fixed ranges
    def widen(ws):
        for row in ws.iter_rows():
            for c in row:
                if isf(c.value):
                    v = c.value
                    v = re.sub(r"('Resource Rate Library'!\$[A-Z]+\$5:\$[A-Z]+\$)85", rf"\g<1>{LIB_LAST}", v)
                    v = re.sub(r"('QS Takeoff'!\$[A-Z]+\$3:\$[A-Z]+\$)(?:194|73|109|137)\b", rf"\g<1>{QS_LAST}", v)
                    c.value = v
    for ws in wb:
        widen(ws)
    for dvx in bu.data_validations.dataValidation:
        if dvx.formula1 and "Resource Rate Library" in dvx.formula1:
            dvx.formula1 = dvx.formula1.replace("$85", f"${LIB_LAST}")
    bu.data_validations.dataValidation = [d for d in bu.data_validations.dataValidation if "X81" not in str(d.sqref)]
    put(bu, "A1", "RESOURCE BUILD-UP - MASTER ESTIMATE WORKBOOK - LINKED COST / DURATION FORMULAS")
    put(bu, "A5", "Weather controls: AK:AM are estimator inputs; AN is formula-controlled exposed Labour/Plant cost feeding Weather Risk. Materials, fixed subcontract and Other costs are excluded from direct weather productivity by default.")
    # ------------------------------------------------------------------ Rate Verification: per-row evidence constants cleared, formulas extended
    rv = wb["Rate Verification"]
    put(rv, "A1", "RATE & ALLOWANCE VERIFICATION - MASTER ESTIMATE WORKBOOK")
    put(rv, "A2", "GREEN requires: (1) current primary evidence, (2) an independent market check, and (3) a first-principles quantity/scope/arithmetic check. Anything less is AMBER; missing or conflicting is RED.")
    rv_formula_cols = [c for c in range(1, 23) if isf(rv.cell(136, c).value)]
    for r in range(8, RV_LAST + 2):
        for c in range(1, 23):
            cell = rv.cell(r, c)
            if isinstance(cell, MergedCell):
                continue
            if r > 136:
                style_like(cell, rv.cell(136, c))
            if c in rv_formula_cols:
                cell.value = tr(rv.cell(136, c).value, rv.cell(136, c).coordinate, cell.coordinate) if r != 136 else cell.value
            elif c == 1:
                cell.value = f"RV-{r - 7:03d}"
            else:
                cell.value = None
    for row in rv.iter_rows(min_row=3, max_row=6):
        for c in row:
            if isf(c.value):
                c.value = c.value.replace("$136", f"${RV_LAST + 1}")
    # ------------------------------------------------------------------ Programme Export: one row per QS row, plus 20 control-task rows
    pe = wb["Programme Export"]
    src_formulas = {c: pe.cell(8, c).value for c in range(9, 25) if isf(pe.cell(8, c).value)}
    plain = {22: pe.cell(8, 22).value, 23: pe.cell(8, 23).value}
    style_src = {c: pe.cell(8, c) for c in range(1, 25)}
    wipe(pe, 7, 90, 1, 24, formulas=True)
    for r in range(7, PE_LAST + 21):
        q = r - 4
        for c in range(1, 25):
            style_like(pe.cell(r, c), style_src[c])
        pe.cell(r, 1).value = r - 6
        if r <= PE_LAST:
            pe.cell(r, 2).value = q
            for c, col in zip(range(3, 9), "ABCEDG"):
                pe.cell(r, c).value = f"=IF('QS Takeoff'!$B{q}=\"\",\"\",'QS Takeoff'!{col}{q})"
        for c, f in src_formulas.items():
            pe.cell(r, c).value = tr(f, pe.cell(8, c).coordinate, pe.cell(r, c).coordinate)
        pe.cell(r, 20).value = f'=IF($D{r}="","",I{r}-N(H{r}))'
        pe.cell(r, 21).value = f'=IF($D{r}="","",IF(ABS(T{r})<0.01,"OK","CHECK"))'
        for c, v in plain.items():
            pe.cell(r, c).value = v
    put(pe, "A1", "PROGRAMME BUILD-UP & MICROSOFT PROJECT EXPORT - MASTER ESTIMATE WORKBOOK")
    put(pe, "A5", "PROGRAMME RECONCILIATION: J takes the largest linked Resource Build-Up applied duration for each QS code; K converts hours to working days; M uses the override in L when one is entered. Rows 7-206 follow QS Takeoff rows 3-202. Control tasks that are not QS items (cure, inspection holds, proof rolls) go in rows 207-226: type the code in D, the name in E, the days in L and predecessor codes in P.")
    # ------------------------------------------------------------------ Risk Truth
    rt = wb["Risk Truth"]
    row53 = {c: rt.cell(53, c).value for c in range(1, 31) if isf(rt.cell(53, c).value)}
    wipe(rt, 25, 40, 1, 30)
    put(rt, "A40", "Exposure ID")
    wipe(rt, 41, 61, 1, 30, formulas=True)
    for r in range(41, 62):
        for c, f in row53.items():
            rt.cell(r, c).value = tr(f, rt.cell(53, c).coordinate, rt.cell(r, c).coordinate)
    for a in ("B5", "B7", "B8"):
        put(rt, a, 0)
    put(rt, "B19", rt["B19"].value.replace("$U$77", f"$U${PE_LAST + 20}") if isf(rt["B19"].value) else rt["B19"].value)
    put(rt, "A1", "RISK TRUTH ENGINE - MASTER ESTIMATE WORKBOOK")
    # ------------------------------------------------------------------ Weather Risk
    wr = wb["Weather Risk"]
    put(wr, "B5", "='Cover & Instructions'!E38")
    for a in ("B6", "B7", "B8", "B9"):
        put(wr, a, None)
    put(wr, "B13", "=0")
    put(wr, "B14", "=0")
    put(wr, "B16", 0)
    for r in range(19, 32):
        put(wr, f"B{r}", None)
        if not isf(wr[f"E{r}"].value):
            put(wr, f"E{r}", None)
    put(wr, "B19", "NZ_Weather_Risk_Estimating_Tool_Rev2.xlsx")
    put(wr, "B29", "No")
    for r in range(33, 40):
        if not isf(wr[f"B{r}"].value):
            put(wr, f"B{r}", None)
    for r in range(42, 49):
        if not isf(wr[f"B{r}"].value):
            put(wr, f"B{r}", None)
        if not isf(wr[f"E{r}"].value):
            put(wr, f"E{r}", None)
    for r in range(51, 58):
        put(wr, f"B{r}", None)
    put(wr, "E9", "Daily time-related preliminaries come from the Build-Up row linked in B13; weekly billed hire from the rows linked in B14.")
    put(wr, "A18", "WEATHER RUN RECORD (from the approved weather tool, if a run has been made)")
    put(wr, "D30", None); put(wr, "E30", None); put(wr, "D31", None); put(wr, "E31", None)
    put(wr, "J53", "Weather Task Type - approved tool library")
    wipe(wr, 74, 83, 1, 8)
    for r in range(86, 92):
        put(wr, f"B{r}", None)
    for r in range(1, wr.max_row + 1):                                    # job-specific explanatory notes on the right-hand panel
        for col in "EF":
            v = wr[f"{col}{r}"].value
            if isinstance(v, str) and not isf(v) and re.search(r"Rev\d\d|Hobsonville|Siteworks|Board", v) and r not in range(5, 17):
                wr[f"{col}{r}"].value = None
    # ------------------------------------------------------------------ Supplier quote register / source & markup
    sq = wb["Supplier Quote Register"]
    put(sq, "A1", "SUPPLIER QUOTE REGISTER")
    put(sq, "A2", "One row per quote or quote request (Q-01, Q-02 ...). Keep consistent with the Quote Log sheet. A quote is evidence, not automatically the controlling amount.")
    wipe(sq, 5, 60, 1, 12)
    sm = wb["Source & Markup Register"]
    put(sm, "A1", "SOURCE & MARKUP REGISTER - PROVENANCE OF QUANTITIES")
    wipe(sm, 3, 40, 1, 12)
    # ------------------------------------------------------------------ Commercial Inputs
    ci = wb["Commercial Inputs"]
    for r in range(5, 25):
        ref = ci.cell(r, 1).value
        if ref in ("CA01", "CA02"):
            ci.cell(r, 6).value = "Yes"
        elif ref in ("CA10",):
            ci.cell(r, 5).value = None
            ci.cell(r, 6).value = "Yes"
        elif ref in ("CA17", "CA18"):
            pass                                                           # company commercial settings are kept as Rick's defaults: set per job
        elif ref == "CA16":
            ci.cell(r, 3).value, ci.cell(r, 4).value, ci.cell(r, 5).value, ci.cell(r, 6).value = "Fixed Amount", None, None, "No"
        elif ref:
            ci.cell(r, 4).value, ci.cell(r, 5).value, ci.cell(r, 6).value = None, None, "No"
    for r in range(5, 30):
        for col in (12, 13):
            if not isf(ci.cell(r, col).value):
                ci.cell(r, col).value = None
    put(ci, "L14", "='Weather Risk'!$B$40")
    put(ci, "M14", "Weather tool run recorded on Weather Risk (CA10 engine feed)")
    notes = {"J8": "Use only after supplier quote returns are reconciled and a distinct uncovered market gap is evidenced.", "I14": "Weather Risk engine feeds this line when the method is Weather Engine.",
             "J14": "Rerun the weather tool when the start date, programme or weather-exposed scope changes.", "I15": "Do not use for weather delay already treated through CA10.",
             "I21": "Company overhead recovery, applied as markup to the subtotal before profit. Set per job.", "I22": "Contractor profit as true margin after overhead and other selected allowances. Set per job."}
    for a, v in notes.items():
        put(ci, a, v)
    wipe(ci, 32, 39, 1, 13)
    put(ci, "A32", "MANAGEMENT RESERVE - ADJUDICATION BASIS")
    for i, col in enumerate("EFGH", 1):
        put(ci, f"{col}33", f"Unresolved item {i} RED")
    for r, lab in zip(range(33, 40), ["Control", "Reserve base", "Selected reserve rate", "Pre-OH/profit Management Reserve", "Tender effect of reserve", "Probability note", "Consumption rule"]):
        put(ci, f"A{r}", lab)
    for col, h in zip("BCD", ["Value", "Classification", "Basis / meaning"]):
        put(ci, f"{col}33", h)
    put(ci, "B38", "N/A"); put(ci, "B39", "N/A")
    put(ci, "D38", "Do not weight any item with a probability unless it is evidenced. Unknown probability = sensitivity only, never an expected value.")
    put(ci, "D39", "Do not automatically consume the reserve when a risk resolves; Rick adjudicates any change under change control.")
    # ------------------------------------------------------------------ Tender Build-Up
    tb = wb["Tender Build-Up"]
    put(tb, "A1", "MASTER ESTIMATE WORKBOOK - TENDER BUILD-UP")
    for a in ("E42", "E43", "F38"):
        put(tb, a, None)
    put(tb, "B3", "='Cover & Instructions'!B38")
    put(tb, "B4", "='Cover & Instructions'!B42")
    put(tb, "A9", "Base direct estimate (by QS section)")
    for r in range(10, 20):
        tb[f"A{r}"].value = None
        tb[f"B{r}"].value = f"=IF(A{r}=\"\",0,SUMIF('QS Takeoff'!$A$3:$A${QS_LAST},A{r},'QS Takeoff'!$G$3:$G${QS_LAST}))"
        tb[f"C{r}"].value = "QS Takeoff (type the section name exactly as used on QS Takeoff)"
    put(tb, "B20", f"=SUM('QS Takeoff'!$G$3:$G${QS_LAST})")
    put(tb, "A21", "Check: QS sections not listed above (must be zero)")
    put(tb, "B21", "=ROUND(B20-SUM(B10:B19),2)")
    put(tb, "C21", '=IF(B21=0,"OK","CHECK - list every QS section above")')
    style_like(tb["A21"], tb["A20"]); style_like(tb["B21"], tb["B20"]); style_like(tb["C21"], tb["C20"])
    wipe(tb, 58, 106, 1, 8, formulas=True)
    # ------------------------------------------------------------------ Consolidated Package (rebuilt generic)
    cp = wb["Consolidated Package"]
    unmerge_all(cp)
    hdr_style, lab_style, num_style = copy.copy(cp["A4"]._style), copy.copy(cp["A5"]._style), copy.copy(cp["B5"]._style)
    wipe(cp, 1, 70, 1, 12, formulas=True)
    S = lambda a, v, st=None: (put(cp, a, v), st and setattr(cp[a], "_style", copy.copy(st)))     # noqa: E731
    S("A1", "CONSOLIDATED ALL-IN PACKAGE - OWN WORKS + SUBCONTRACT PACKAGES")
    S("A2", "WORKING ALL-IN TENDER EXCL GST", lab_style); S("B2", "=F56", num_style)
    for col, h in zip("ABCDEF", ["Package / Control", "Amount excl GST", "Classification", "Included in Siteworks direct?", "Status", "Release"]):
        S(f"{col}4", h, hdr_style)
    S("A5", "Own works current pre-commercial package", lab_style); S("B5", "='Tender Build-Up'!B40", num_style); S("C5", "LINKED / CURRENT"); S("D5", "Yes"); S("E5", "Own works"); S("F5", "AMBER / WORKING")
    for r in range(6, 12):
        S(f"A{r}", None, lab_style); S(f"B{r}", None, num_style); S(f"C{r}", "QUOTED"); S(f"D{r}", "No"); S(f"F{r}", "AMBER")
    S("A12", "GROSS CURRENT KNOWN PACKAGE", lab_style); S("B12", "=B5+SUM(B6:B11)", num_style); S("C12", "CALCULATED"); S("D12", "Review total only"); S("E12", "Not final tender"); S("F12", "HOLD")
    S("A13", "Confirmed overlap between own works and a package", lab_style); S("B13", 0, num_style); S("C13", "CLOSED"); S("D13", "None"); S("E13", "Interface closed"); S("F13", "AMBER / CLOSED")
    S("A14", "Current known package after overlap", lab_style); S("B14", "=B12-B13", num_style); S("C14", "CALCULATED"); S("D14", "Review only"); S("E14", "Not final tender"); S("F14", "HOLD")
    S("A15", "Unpriced packages (see Quote Log)", lab_style); S("C15", "UNPRICED"); S("D15", "No")
    S("A16", "WORKING ALL-IN TENDER EXCL GST", lab_style); S("B16", "=F56", num_style); S("C16", "CALCULATED / WORKING"); S("D16", "N/A")
    S("A20", "Rules: one row per subcontract package (rows 6-11; insert more by copying a row). Amount as a number, 'Included in Siteworks direct?' = No, classification QUOTED / QUOTED BASE / CONFIRMED ADDITIONAL / ALLOCATED (never CALCULATED). A package with no price yet stays as a row with no amount and is recorded on the Quote Log.")
    for col, h in zip("ABCDEFG", ["Scenario", "Precommercial cost", "Overhead", "Subtotal after OH", "Profit", "Tender excl GST", "Commercial uplift"]):
        S(f"{col}55", h, hdr_style)
    S("A54", "CURRENT WORKING COMMERCIAL SCENARIO - OVERHEAD AND PROFIT FROM COMMERCIAL INPUTS CA17 / CA18")
    S("A56", "Current known package after overlap", lab_style); S("B56", "=B14", num_style)
    S("C56", "=IF('Commercial Inputs'!$F$21<>\"Yes\",0,IF('Commercial Inputs'!$C$21=\"Markup %\",B56*N('Commercial Inputs'!$D$21),N('Commercial Inputs'!$E$21)))", num_style)
    S("D56", "=B56+C56", num_style)
    S("E56", "=IF('Commercial Inputs'!$F$22<>\"Yes\",0,IF('Commercial Inputs'!$C$22=\"Margin %\",D56/(1-N('Commercial Inputs'!$D$22))-D56,IF('Commercial Inputs'!$C$22=\"Markup %\",D56*N('Commercial Inputs'!$D$22),N('Commercial Inputs'!$E$22))))", num_style)
    S("F56", "=D56+E56", num_style); S("G56", "=F56-B56", num_style)
    # ------------------------------------------------------------------ Cover & Instructions (rebuilt generic)
    build_cover(wb)
    # ------------------------------------------------------------------ logs and guide
    add_logs(wb)
    add_guide(wb)
    order = ["Template Guide"] + KEEP + ["Email Log", "RFI Log", "Quote Log", "Document Register", "Decision Log", "Contacts", "Change Log"]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    wb.save(out)


def build_cover(wb) -> None:
    ws = wb["Cover & Instructions"]
    keep_styles = {k: copy.copy(ws[k]._style) for k in ("A1", "A2", "A4", "E4", "A5", "A9", "E5", "F5", "A21", "A22", "B22", "C22", "A33", "A34", "A38", "B38", "D38", "A46", "A47", "B47",
                                                         "A48", "B48", "C48", "A57", "A58", "B58", "A70", "A71", "B71", "A87", "A88", "A89", "B89", "A100", "A101", "A102", "B102", "C102", "A130", "A131", "B131")}
    unmerge_all(ws)
    wipe(ws, 1, 200, 1, 10, formulas=True)
    ST = lambda a, k: setattr(ws[a], "_style", copy.copy(keep_styles[k]))      # noqa: E731

    def w(a, v, k=None):
        ws[a].value = v
        if k:
            ST(a, k)
    w("A1", "TENDER SNAPSHOT / START HERE", "A1")
    ws.merge_cells("A1:H1")
    w("A2", '="Working all-in tender excl. GST: "&TEXT(\'Consolidated Package\'!F56,"$#,##0.00")&" | overhead "&TEXT(\'Commercial Inputs\'!D21,"0.0%")&" + profit "&TEXT(\'Commercial Inputs\'!D22,"0.0%")', "A2")
    ws.merge_cells("A2:H2")
    w("A4", "THE FIRST QUESTION: WHAT WILL THIS COST?", "A4"); ws.merge_cells("A4:C4")
    w("E4", "CURRENT COMMERCIAL BUILD-UP", "E4"); ws.merge_cells("E4:H4")
    w("A5", "='Consolidated Package'!F56", "A5"); ws.merge_cells("A5:C8")
    rows = [("Own works pre-commercial", "='Consolidated Package'!B5"), ("Subcontract packages (known)", "=SUM('Consolidated Package'!B6:B11)"), ("Current known cost before OH/profit", "='Consolidated Package'!B56"),
            ("Company overhead", "='Consolidated Package'!C56"), ("Profit", "='Consolidated Package'!E56"), ("WORKING ALL-IN TENDER", "='Consolidated Package'!F56"),
            ("Gross margin on sell", "=IFERROR(('Consolidated Package'!F56-'Consolidated Package'!B56)/'Consolidated Package'!F56,0)")]
    for i, (lab, f) in enumerate(rows):
        w(f"E{5 + i}", lab, "E5"); w(f"F{5 + i}", f, "F5")
    w("A9", "WORKING ALL-IN TENDER, EXCL. GST", "A9"); ws.merge_cells("A9:C9")
    w("E13", "KEY CONTROLS", "E4"); ws.merge_cells("E13:H13")
    ctl = [("Own works standalone tender", "='Tender Build-Up'!B47"), ("Management Reserve, pre-margin", "='Commercial Inputs'!E20"), ("Weather allowance", "='Weather Risk'!B38"),
           ("Dry programme (working days)", "='Weather Risk'!B8"), ("Weather-adjusted planning (days)", "='Weather Risk'!B8+'Weather Risk'!B26"), ("Risk Truth release gate", "='Risk Truth'!B15")]
    for i, (lab, f) in enumerate(ctl):
        w(f"E{14 + i}", lab, "E5"); w(f"F{14 + i}", f, "F5")
    w("A21", "HOW TO READ THIS WORKBOOK - SIMPLE WORKFLOW", "A21"); ws.merge_cells("A21:H21")
    flow = [("Cover & Instructions", "Job details, document register, assumptions, exclusions, RFIs, release QA"), ("Consolidated Package", "Own works + subcontract packages + overhead + profit = all-in tender"),
            ("QS Takeoff", "One row per priced item with quantity, source, status"), ("Resource Build-Up", "How each QS item is built: labour, plant, materials, subcontract, preliminaries"),
            ("Commercial Inputs", "Allowances, management reserve, overhead, profit, GST"), ("Tender Build-Up", "The commercial stack, calculated"), ("Risk Truth", "Risk register, exposures, duplication controls, release gate"),
            ("Weather Risk / Programme Export", "Project weather inputs and task durations / logic"), ("Supplier Quote Register + logs", "Quotes, emails, RFIs, documents, decisions, changes")]
    for i, (a, b) in enumerate(flow):
        w(f"A{22 + i}", i + 1, "A22"); w(f"B{22 + i}", a, "B22"); w(f"C{22 + i}", b, "C22"); ws.merge_cells(f"C{22 + i}:H{22 + i}")
    w("A33", "CURRENT RELEASE POSITION", "A33"); ws.merge_cells("A33:H33")
    w("A34", "(Write where the estimate stands and what blocks release. List every RED item.)", "A34"); ws.merge_cells("A34:H36")
    labels = [("A38", "Project", "B38", "D38", "Location", "E38"), ("A39", "Current drawings", "B39", "D39", "Issue date", "E39"), ("A40", "Scope", "B40", "D40", "GST", "E40"),
              ("A41", "Programme", "B41", "D41", "Weather", "E41"), ("A42", "Price status", "B42", "D42", "Commercial", "E42")]
    for la, ltxt, va, lb, ltxt2, vb in labels:
        w(la, ltxt, "A38"); ST(va, "B38"); w(lb, ltxt2, "D38"); ST(vb, "B38")
    for a in ("B38", "E38", "B39", "E39", "B40", "B41", "E41", "B42", "E42"):
        ws[a].value = "NOT PROVIDED"
    w("E40", "All values shown exclude GST")
    w("A44", "Risk Truth", "A38"); w("B44", "Risk-control and reconciliation module. Tracks baseline vs residual exposure, provision routing, duplication and release gate.", "B38")
    # ---- document register
    w("A46", "DRAWING / DOCUMENT REGISTER", "A46"); ws.merge_cells("A46:H46")
    for col, h in zip("ABC", ["Document / Drawing", "Reference", "Estimate use"]):
        w(f"{col}47", h, "A47")
    cats = ["Plans / drawings", "Specifications", "Geotechnical / ground information", "Survey / CAD / models", "Tender / contract / scope documents", "Supplier / subcontract quotations",
            "Other consultants / approvals", "Correspondence / addenda / clarifications", "Backend construction intelligence"]
    for i, c in enumerate(cats):
        w(f"A{48 + i}", c, "A48"); ST(f"B{48 + i}", "B48"); ST(f"C{48 + i}", "C48")
    # ---- assumptions / exclusions
    r = 57
    w(f"A{r}", "ASSUMPTIONS / QUALIFICATIONS / RFIs", "A57"); ws.merge_cells(f"A{r}:H{r}")
    for i in range(1, 16):
        r += 1; w(f"A{r}", f"A{i}", "A58"); ST(f"B{r}", "B58")
    for i in range(1, 9):
        r += 1; w(f"A{r}", f"EX-{i:02d}", "A58"); ST(f"B{r}", "B58")
    r += 1
    w(f"A{r}", "PRICING QUALIFICATIONS / RFIs", "A70"); ws.merge_cells(f"A{r}:H{r}")
    for i in range(1, 31):
        r += 1; w(f"A{r}", f"RFI-{i:02d}", "A71"); ST(f"B{r}", "B71")
    r += 1
    w(f"A{r}", "KEY RISKS / WATCHOUTS", "A87"); ws.merge_cells(f"A{r}:H{r}")
    r += 1
    for col, h in zip("ABCD", ["Risk ID", "Risk", "Likelihood", "Impact"]):
        w(f"{col}{r}", h, "A88")
    for k in range(25, 40):
        r += 1
        for col, rc in zip("ABCD", "ABDE"):
            w(f"{col}{r}", f"=IF('Risk Truth'!{rc}{k}=\"\",\"\",'Risk Truth'!{rc}{k})", "A89" if col == "A" else "B89")
    r += 2
    w(f"A{r}", "TENDER RELEASE QA", "A100"); ws.merge_cells(f"A{r}:H{r}")
    r += 1
    for col, h in zip("ABCDE", ["Section", "Requirement", "Status (colour + reason)", "Owner", "Required Before Tender?"]):
        w(f"{col}{r}", h, "A101")
    qa = [("References", "Current Intake, Master Prompt, Triple Verification and applicable standards recorded"), ("Documents", "Current drawing/specification/tender register complete"),
          ("Scope", "Work-package inclusions, exclusions and interfaces defined"), ("Takeoff", "Major quantities measured/provided/calculated and referenced to current documents"),
          ("Resource build-up", "Every active QS code has linked labour/plant/material/subcontract rows or an explicit PC"), ("Rate verification", "Important rates classified and three-way verified or flagged"),
          ("Quotes / PCs", "PC sums and quote-required packages reconciled to the Quote Log"), ("Assumptions", "Assumptions and exclusions are explicit, current and defensible"),
          ("Interfaces", "By-others / client / authority / consultant responsibilities identified"), ("Risk Truth", "All material risks have baseline, residual exposure and cost treatment"),
          ("Risk duplication", "No residual risk is costed twice; overlap groups and cost-treatment owners checked"), ("Weather", "Weather Risk reviewed; approved tool run current or documented"),
          ("Ground / geotech", "Ground and groundwater risks aligned to current geotechnical information"), ("Earthworks / haulage", "Balance, bulking, haul, disposal and import assumptions defined"),
          ("Temporary works", "Temporary works requirements and responsibility identified"), ("Programme", "Working days, productive crew days and elapsed duration reconciled"),
          ("Programme cost check", "Programme Export resource/QS cost checks all read OK"), ("Commercial", "Overhead, profit, escalation, securities and commercial inputs confirmed"),
          ("Pricing schedule", "Client pricing schedule mapped without double counting"), ("Independent review", "Review completed or consciously waived by the estimator"), ("Release", "Risk Truth, Weather and commercial gates reviewed"),
          ("Approval", "Final tender/estimate approved by the authorised reviewer")]
    for sec, req in qa:
        r += 1; w(f"A{r}", sec, "A102"); w(f"B{r}", req, "B102"); ST(f"C{r}", "C102")
    r += 2
    w(f"A{r}", "TRIPLE VERIFICATION & RISK TRUTH - MANDATORY", "A130"); ws.merge_cells(f"A{r}:H{r}")
    rules = ["No rate, quantity, productivity, duration, allowance or risk treatment is GREEN without current primary evidence, an independent check and a first-principles check.",
             "Quantifiable physical risk consequences return to Resource Build-Up / QS; only distinct residual uncertainty goes to Risk Truth.",
             "CA01 and CA02 default to Risk Truth residual feeds; manual percentages are not used unless documented.",
             "Weather remains in the Weather Risk specialist path and is never also added inside rates.",
             "Any RED duplication, programme reconciliation or unresolved material RED exposure blocks final release."]
    for i, t in enumerate(rules, 1):
        r += 1; w(f"A{r}", i, "A131"); w(f"B{r}", t, "B131")
    r += 2
    w(f"A{r}", "REVISION RECORD (full history on the Change Log sheet)", "A130"); ws.merge_cells(f"A{r}:H{r}")
    for col, h in zip("ABC", ["Revision", "Date", "Description"]):
        r_ = r + 1; w(f"{col}{r_}", h, "A101")
    for k in range(3):
        for col in "ABC":
            ST(f"{col}{r + 2 + k}", "B102")


def hdr(ws, row, headers, widths=None, style_src=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        if style_src is not None:
            style_like(c, style_src)
        else:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F1F1F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[c.column_letter].width = widths[i - 1] if i - 1 < len(widths) else 18
    ws.freeze_panes = ws.cell(row + 1, 2)


LOGS = {
    "Email Log": (["Email ID", "Date / time", "Direction (In/Out)", "From", "To", "Cc", "Subject", "Source reference (thread/message ID or file)", "Category", "Linked refs (RFI / Quote / QS codes)", "Summary (1-3 neutral sentences)",
                   "Key wording (exact words that matter)", "Attachments (and Document Register IDs)", "Action required (Yes/No + what)", "Action owner", "Due", "Status (Open/Closed/Info only)",
                   "Estimate impact (None/Quantity/Rate/Scope/Programme/Commercial)", "Applied in estimate (revision + line, or Not yet)"],
                  [11, 17, 12, 24, 24, 20, 38, 30, 16, 24, 54, 44, 32, 28, 16, 12, 14, 24, 28]),
    "RFI Log": (["RFI ID", "Date raised", "Raised by", "Addressed to", "Subject", "Question - what is missing", "Where expected (document / drawing)", "Why it matters", "Decision impact (cost / programme / scope)",
                 "QS codes / risk IDs affected", "Assumption in force until answered", "Status (OPEN/CLOSED/SUPERSEDED)", "RAG", "Date answered", "Answer summary", "Answer source (Email ID / document)",
                 "Resulting estimate change (revision + line)", "Owner"],
                [10, 12, 14, 18, 32, 54, 30, 40, 40, 22, 44, 16, 8, 12, 44, 24, 30, 14]),
    "Quote Log": (["Quote ID", "Package / scope", "Supplier / subcontractor", "Contact (name only)", "Quote number", "Quote date", "Valid until", "Source (Email ID or file)", "Scope included", "Exclusions and qualifications",
                   "Price excl. GST", "GST basis", "Unit rates (if given)", "Payment / retention terms", "Lead time", "QS codes fed", "Treatment (Overlay / Fed into library rate / Reference only)",
                   "Amount used in estimate", "Why this amount", "Comparison with other quotes", "Evidence strength (GREEN/AMBER/RED)", "Status", "Next action / due", "Notes"],
                  [9, 28, 24, 18, 16, 12, 12, 26, 44, 44, 14, 12, 26, 22, 12, 18, 30, 14, 36, 36, 14, 22, 24, 30]),
    "Document Register": (["Doc ID", "Title", "Type (Plan/Specification/Schedule/Geotech/Consent/Addendum/Photo/Email/Other)", "Reference / drawing number", "Revision", "Date", "Status (CURRENT/SUPERSEDED/UNCLEAR)",
                           "Superseded by", "Source (file name, link, Email ID)", "Received date", "Used for (QS codes)", "Notes"],
                          [10, 40, 30, 22, 10, 12, 18, 16, 36, 12, 22, 40]),
    "Decision Log": (["Decision ID", "Date", "Decision", "Made by", "Options considered", "Basis / evidence", "Estimate impact", "Revision applied", "Status (Proposed/Approved by estimator/Superseded)"],
                     [11, 12, 50, 16, 40, 44, 30, 14, 24]),
    "Contacts": (["Organisation", "Person", "Role", "Package / trade", "Source of contact (Email ID)", "Notes (no phone numbers or home addresses)"], [28, 22, 22, 22, 24, 44]),
    "Change Log": (["Revision", "Date", "What changed (scope / quantity / rate / quote / commercial / risk / programme)", "Reason (document, email, decision)", "Lines affected", "Cost movement excl. GST (from the workbook, not estimated)", "Recorded by"],
                   [10, 12, 54, 40, 26, 22, 16]),
}


def add_logs(wb) -> None:
    donor = wb["Resource Rate Library"]["A4"]
    for name, (heads, widths) in LOGS.items():
        ws = wb.create_sheet(name)
        hdr(ws, 1, heads, widths, donor)
        ws.row_dimensions[1].height = 45
        for c in range(1, len(heads) + 1):
            ws.cell(1, c).alignment = Alignment(wrap_text=True, vertical="center")


def add_guide(wb) -> None:
    ws = wb.create_sheet("Template Guide")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    rows = [
        ("MASTER ESTIMATE WORKBOOK - TEMPLATE", f"Blackline QSi Pro  |  template {VERSION}"),
        ("Purpose", "Clean master for one job. An external AI (or the estimator) fills it from the tender documents; the Blackline QSi Pro app imports it and recalculates every figure itself."),
        ("Instructions", "Read instructions/Workbook_Builder_Instructions.md before filling anything. The sheet names, header rows and column positions below are a contract with the app."),
        ("", ""),
        ("CAPACITY", ""),
        ("QS Takeoff", f"Rows 3-{QS_LAST} (one item per row; total on row {QS_LAST + 1}). Formula columns F, G, L-O, R are pre-filled."),
        ("Resource Build-Up", f"Rows 7-{BU_LAST}. All formula columns are pre-filled; type only the input columns (see instructions section 5.2)."),
        ("Resource Rate Library", f"Rows 5-{LIB_LAST}. The existing company rates are kept; add rates only with source, date and verification status."),
        ("Programme Export", f"Rows 7-{PE_LAST} follow QS rows 3-{QS_LAST}; rows {PE_LAST + 1}-{PE_LAST + 20} are for control tasks (cure, holds) that are not QS items."),
        ("Rate Verification", f"One row per Build-Up row (rows 8-{RV_LAST + 1}); fill the evidence columns K-P, T, U."),
        ("Tender Build-Up", "Type the QS section names in A10:A19 (exactly as written on QS Takeoff). Row 21 must read OK."),
        ("Cover & Instructions", "Assumptions A1-A15, exclusions EX-01 to EX-08, RFIs RFI-01 to RFI-30, 9 document categories. Copy a row inside a block to add more."),
        ("Consolidated Package", "Rows 6-11 are subcontract package slots. Row 13 holds a confirmed overlap."),
        ("Weather Risk", "B13 and B14 are set to =0 in the template. Point them at the Build-Up rows that carry daily time-related preliminaries and weekly hire, e.g. B13 ='Resource Build-Up'!Y22*8 and B14 ='Resource Build-Up'!Y48+'Resource Build-Up'!Y60."),
        ("", ""),
        ("COMPANY DEFAULTS KEPT", "Overhead 10% (markup) and profit 15% (true margin) on Commercial Inputs CA17/CA18 are the company's current settings; confirm per job. All other allowances are OFF and zero."),
        ("NOT IN THIS TEMPLATE", "Summary, Stage Cashflow, Estimate Review Board, Materials & Component Takeoff and MS Project Import sheets: the platform produces summaries and the programme export itself."),
        ("LOGS", "Email Log, RFI Log, Quote Log, Document Register, Decision Log, Contacts, Change Log: headers are fixed (instructions section 8)."),
        ("RECALCULATE", "After editing, recalculate and save so formulas hold values (LibreOffice: soffice --headless --convert-to xlsx --outdir out file.xlsx)."),
    ]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
    print("written", sys.argv[2])
