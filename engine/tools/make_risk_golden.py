"""Build tests/golden_risk_truth.json: random Risk Truth scenarios run through the REAL Master Estimate Workbook's own Risk Truth formulas
(Rev34E 'Risk Truth' tab), recalculated by LibreOffice, so tender_engine.risk_truth can be tested against the workbook's answers.

Run:  python tools/make_risk_golden.py private/Rev34E.xlsx [n_scenarios]
Needs LibreOffice (soffice). The fixture holds only synthetic inputs and the workbook's computed numbers (no client data).
The workbook rows 41-61 are replaced by the scenario, using the template formulas from the workbook's own spare row 53.
"""
from __future__ import annotations
import json
import random
import re
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from tender_engine.risk_truth import MODES  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "tests" / "golden_risk_truth.json"
COL = {"ref": "A", "risk_id": "B", "consequence": "C", "driver": "D", "cost_ref": "E", "treatment": "F", "baseline_amount": "G", "residual": "H", "mode": "I",
       "qty": "J", "unit": "K", "unit_rate": "L", "probability": "M", "provision": "O", "owner": "P", "distinct": "Q", "target": "R", "include": "S",
       "overlap_group": "U", "embedded": "V", "evidence": "W", "integration_complete": "Y"}
FIRST, LAST, TEMPLATE = 41, 61, 53
REFS = ["EW01", "ew01", "PV01", "RW02", "ES03", "CA10", "PG01 / Programme", "EW08/EW09", "", "QS-X"]


def scenario(rng: random.Random, i: int) -> dict:
    n = rng.randint(0, 14) if i else 0
    items = []
    for j in range(n):
        mode = rng.choice(MODES)
        it = {"ref": f"T{i}-{j}", "risk_id": f"RT{rng.randint(1, 9):02d}", "cost_ref": rng.choice(REFS), "mode": mode,
              "qty": rng.choice(["", "", 10, 250.5]), "unit_rate": rng.choice(["", 40, 1234.56]),
              "probability": rng.choice(["", 0.3, 0.65, 1]), "provision": rng.choice(["", 0, 5000, 12345.67]),
              "owner": rng.choice(["Primary", "Contributing", "Reference Only"]), "distinct": rng.choice(["Yes", "Yes", "No", "N/A"]),
              "target": rng.choice(["CA01", "CA02", "ca01", "Direct QS", "CA02"]), "include": rng.choice(["Yes", "Yes", "No"]),
              "overlap_group": rng.choice(["", "", "OG-1"]), "embedded": rng.choice(["", "No", "CHECK - confirm", "Yes"]),
              "evidence": rng.choice(["RED", "AMBER", "GREEN", "GREEN"]), "integration_complete": rng.choice(["Yes", "No"])}
        items.append(it)
    risks = [rng.choice(["RED", "AMBER", "GREEN"]) for _ in range(rng.randint(0, 15))]
    if i < 4:                      # clean scenarios so the gate is seen to PASS: no RED risks, and (from scenario 2) tidy green exposures
        risks = [r if r != "RED" else "AMBER" for r in risks]
        items = [{"ref": f"G{i}-{j}", "risk_id": "RT01", "cost_ref": f"CR{j}", "mode": "Fixed Allowance", "provision": 1000 * (j + 1), "owner": "Primary",
                  "distinct": "Yes", "target": "CA02" if j % 2 else "CA01", "include": "Yes", "evidence": "GREEN", "integration_complete": "Yes"}
                 for j in range(0 if i < 2 else i)]
    return {"items": items, "risk_rags": risks}


def build(src: str, sc: dict, path: Path) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb["Risk Truth"]
    tmpl = {c: ws[f"{c}{TEMPLATE}"].value for c in ("N", "T", "X", "Z")}
    assert all(isinstance(v, str) and v.startswith("=") for v in tmpl.values()), "template formulas not found in row 53"
    for r in range(FIRST, LAST + 1):
        for c in range(1, 30):
            ws.cell(r, c).value = None
        for c, f in tmpl.items():
            ws[f"{c}{r}"] = re.sub(r"(?<=[A-Z])53\b", str(r), f)
    for j, it in enumerate(sc["items"]):
        for k, c in COL.items():
            v = it.get(k, "")
            if v != "":
                ws[f"{c}{FIRST + j}"] = v
    for j in range(15):
        ws[f"K{25 + j}"] = sc["risk_rags"][j] if j < len(sc["risk_rags"]) else None
    wb.save(path)


def main(src: str, n: int) -> None:
    rng = random.Random(20261014)
    scs = [scenario(rng, i) for i in range(n)]
    tmp = Path(tempfile.mkdtemp(prefix="rtgold_"))
    for i, sc in enumerate(scs):
        build(src, sc, tmp / f"s{i}.xlsx")
    outdir = tmp / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), *[str(tmp / f"s{i}.xlsx") for i in range(n)]],
                   check=True, capture_output=True, timeout=3000)
    for i, sc in enumerate(scs):
        v = openpyxl.load_workbook(outdir / f"s{i}.xlsx", data_only=True)["Risk Truth"]
        rows = []
        for j in range(len(sc["items"])):
            r = FIRST + j
            rows.append({k: v[f"{c}{r}"].value for k, c in (("modelled", "N"), ("addition", "T"), ("duplication", "X"), ("status", "Z"))})
        sc["expected"] = {"rows": rows, "ca01": v["B9"].value, "ca02": v["B10"].value, "red_risks": v["B16"].value, "red_exposures": v["B17"].value,
                          "duplicate_gate": v["B18"].value, "programme_gate": v["B19"].value, "gate": v["B15"].value}
    OUT.write_text(json.dumps({"source": "Rev34E Risk Truth tab formulas recalculated in LibreOffice", "scenarios": scs}, indent=1, default=str) + "\n")
    shutil.rmtree(tmp)
    print("wrote", OUT, len(scs), "scenarios")


if __name__ == "__main__":
    main(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 40)
